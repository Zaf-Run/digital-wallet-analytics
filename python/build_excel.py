"""Builds wallet_analytics_workbook.xlsx from the summary CSVs. Fee model is formula driven."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

O = "../outputs/"
F = Font(name="Arial", size=10)
B = Font(name="Arial", size=10, bold=True)
BLUE = Font(name="Arial", size=10, color="0000FF")
GREEN = Font(name="Arial", size=10, color="008000")
HDR = PatternFill("solid", start_color="1F3864")
HF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
YEL = PatternFill("solid", start_color="FFFF00")
wb = Workbook()

def put_df(ws, df, r0=1, fmts=None):
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(r0, j, c); x.font = HF; x.fill = HDR; x.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(30, len(str(c)) + 3))
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            x = ws.cell(i, j, None if pd.isna(v) else (v.item() if hasattr(v, "item") else v)); x.font = F
            f = (fmts or {}).get(df.columns[j - 1])
            if f: x.number_format = f
    ws.freeze_panes = ws.cell(r0 + 1, 1)

# README
ws = wb.active; ws.title = "README"
lines = ["Digital Wallet Analytics workbook",
 "Source: PaySim mobile money log, 6,362,620 rows, 743 hourly steps (31 days, day 31 has 23 hours). Simulated currency units.",
 "Sheets: Assumptions and Fee_Model are formula driven. KPI, Daily, Type, Hourly, Balance, Segments, Validation hold results exported by python/analysis.py.",
 "Blue cells on Assumptions are inputs. Change them and Fee_Model recalculates.",
 "Revenue is a SCENARIO. The source file has no fee records. Fraud labelled rows earn no fee (assumed reversed).",
 "Fee_Model uses band averages, so it approximates the exact per transaction fee cap. Column 'Exact (Python)' shows the exact base case for comparison.",
 "Balances are pre and post transaction balances on each row, not snapshots. Only 43% of rows reconcile. See Balance and Validation sheets."]
for i, t in enumerate(lines, 1):
    ws.cell(i, 1, t).font = B if i == 1 else F
ws.column_dimensions["A"].width = 140

# Assumptions
fa = pd.read_csv(O + "fee_assumptions.csv")
wa = wb.create_sheet("Assumptions")
wa["A1"] = "Fee schedule assumptions (illustrative, not observed)"; wa["A1"].font = B
for j, c in enumerate(["type", "gross_rate", "partner_rate", "cap_per_txn"], 1):
    x = wa.cell(3, j, c); x.font = HF; x.fill = HDR
for i, r in enumerate(fa.itertuples(index=False), 4):
    wa.cell(i, 1, r.type).font = F
    for j, v in enumerate([r.gross_rate, r.partner_rate, r.cap], 2):
        x = wa.cell(i, j, float(v)); x.font = BLUE; x.number_format = "0.0%" if j < 4 else "#,##0"
wa["A10"] = "Gross rate multiplier"; wa["B10"] = 1.0; wa["B10"].font = BLUE; wa["B10"].fill = YEL
wa["A11"] = "Use 0.5 for low, 1.0 base, 1.5 high. Cap applies per transaction to the gross fee."
wa["A12"] = "Partner rate is the cost paid to agents, processors or banks. Net revenue = gross fee minus partner cost."
for a in ("A10", "A11", "A12"): wa[a].font = F
wa.column_dimensions["A"].width = 28
for c in "BCD": wa.column_dimensions[c].width = 16

# Fee model
fm = pd.read_csv(O + "fee_model_inputs.csv")
wf = wb.create_sheet("Fee_Model")
heads = ["type", "amount_band", "billable_txns", "billable_value", "avg_amount", "gross_rate", "cap", "gross_fee", "partner_cost", "net_fee"]
for j, c in enumerate(heads, 1):
    x = wf.cell(1, j, c); x.font = HF; x.fill = HDR; wf.column_dimensions[get_column_letter(j)].width = 18
n = len(fm)
for i, r in enumerate(fm.itertuples(index=False), 2):
    wf.cell(i, 1, r.type); wf.cell(i, 2, str(r.amount_band)); wf.cell(i, 3, int(r.txn_count)).font = BLUE; wf.cell(i, 4, float(r.value)).font = BLUE
    wf.cell(i, 5, f"=IF(C{i}=0,0,D{i}/C{i})")
    wf.cell(i, 6, f"=INDEX(Assumptions!$B$4:$B$8,MATCH(A{i},Assumptions!$A$4:$A$8,0))*Assumptions!$B$10")
    wf.cell(i, 7, f"=INDEX(Assumptions!$D$4:$D$8,MATCH(A{i},Assumptions!$A$4:$A$8,0))")
    wf.cell(i, 8, f"=C{i}*MIN(E{i}*F{i},G{i})")
    wf.cell(i, 9, f"=IF(INDEX(Assumptions!$B$4:$B$8,MATCH(A{i},Assumptions!$A$4:$A$8,0))=0,0,H{i}*INDEX(Assumptions!$C$4:$C$8,MATCH(A{i},Assumptions!$A$4:$A$8,0))/INDEX(Assumptions!$B$4:$B$8,MATCH(A{i},Assumptions!$A$4:$A$8,0)))")
    wf.cell(i, 10, f"=H{i}-I{i}")
    for j in range(1, 11):
        c = wf.cell(i, j)
        if c.font != BLUE and not (j in (3, 4)): c.font = F if j < 5 else F
    for j in (3, 4, 5, 7, 8, 9, 10): wf.cell(i, j).number_format = "#,##0"
    wf.cell(i, 6).number_format = "0.00%"
last = n + 1
wf.freeze_panes = "A2"
# summary by type
wf["L1"] = "Summary by type"; wf["L1"].font = B
for j, c in enumerate(["type", "gross_fee", "partner_cost", "net_fee", "net_share", "Exact net (Python base)", "Approx error"], 12):
    x = wf.cell(2, j, c); x.font = HF; x.fill = HDR; wf.column_dimensions[get_column_letter(j)].width = 20
ex = pd.read_csv(O + "revenue_scenarios.csv"); ex = ex[(ex.scenario == "base") & (ex.fee_cap_applied)].set_index("type").net_revenue
for k, t in enumerate(["CASH_IN", "CASH_OUT", "DEBIT", "PAYMENT", "TRANSFER"], 3):
    wf.cell(k, 12, t).font = F
    for j, col in zip((13, 14, 15), "HIJ"):
        wf.cell(k, j, f"=SUMIFS(${col}$2:${col}${last},$A$2:$A${last},$L{k})").number_format = "#,##0"
    wf.cell(k, 16, f"=IF($O$8=0,0,O{k}/$O$8)").number_format = "0.0%"
    x = wf.cell(k, 17, float(ex[t])); x.font = BLUE; x.number_format = "#,##0"
    wf.cell(k, 18, f"=IF(Q{k}=0,0,O{k}/Q{k}-1)").number_format = "0.0%"
wf.cell(8, 12, "Total").font = B
for j in (13, 14, 15, 17):
    c = get_column_letter(j); wf.cell(8, j, f"=SUM({c}3:{c}7)").number_format = "#,##0"
wf.cell(8, 18, "=IF(Q8=0,0,O8/Q8-1)").number_format = "0.0%"
wf["L10"] = "Exact column is valid only when the multiplier is 1.0. Errors come from using band averages with a per transaction cap."
wf["L10"].font = F

# result sheets
def sheet(name, csv, fmts=None):
    d = pd.read_csv(O + csv); w = wb.create_sheet(name); put_df(w, d, 1, fmts); return w, d
k = pd.read_csv(O + "kpi_summary.csv"); wk = wb.create_sheet("KPI"); put_df(wk, k, 1, {"value": "#,##0.0000"}); wk.column_dimensions["A"].width = 40; wk.column_dimensions["B"].width = 24
wd, d = sheet("Daily", "daily_summary.csv", {c: "#,##0" for c in ["txn_count", "txn_value", "dau", "gross_revenue", "net_revenue"]})
for col in ("fraud_rate_count", "recon_exception_rate"):
    j = list(d.columns).index(col) + 1
    for i in range(2, len(d) + 2): wd.cell(i, j).number_format = "0.00%"
ch = LineChart(); ch.title = "Daily active users (non fraud senders)"; ch.height = 8; ch.width = 18
jd = list(d.columns).index("dau") + 1
ch.add_data(Reference(wd, min_col=jd, min_row=1, max_row=len(d) + 1), titles_from_data=True)
ch.set_categories(Reference(wd, min_col=1, min_row=2, max_row=len(d) + 1)); wd.add_chart(ch, "A36")
wt, _ = sheet("Type", "type_summary.csv", {"txn_value": "#,##0", "avg_amount": "#,##0", "count_share": "0.0%", "value_share": "0.0%", "fraud_rate_count": "0.00%", "recon_exception_rate": "0.0%"})
sheet("Hourly", "hourly_summary.csv", {"fraud_rate_count": "0.00%"})
sheet("Frequency", "txn_frequency.csv", {"share_of_customers": "0.00%"})
sheet("Segments", "customer_value_segments.csv", {c: "0.0%" for c in ["customers_share", "value_share", "gross_revenue_share", "net_revenue_share"]})
sheet("Balance", "balance_summary.csv", {c: "#,##0" for c in ["mean", "median", "p25", "p75", "p90", "p99"]})
sheet("Fraud_Rules", "fraud_rule_evaluation.csv", {"precision": "0.0%", "recall": "0.0%", "f1": "0.000", "fraud_value_captured_share": "0.0%"})
sheet("Validation", "validation_checks.csv")
wb["Validation"].column_dimensions["A"].width = 60; wb["Validation"].column_dimensions["E"].width = 60
wb.save("wallet_analytics_workbook.xlsx")
